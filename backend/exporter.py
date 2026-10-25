"""HTML draft -> .docx using python-docx."""
import io

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Pt

DISCLAIMER = "AI-assisted draft. Review and verify all content and citations before use."


def _runs(par, node, bold=False, italic=False):
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child).replace("\n", " ")
            if text.strip() or text == " ":
                r = par.add_run(text)
                r.bold, r.italic = bold, italic
        elif child.name in ("strong", "b"):
            _runs(par, child, True, italic)
        elif child.name in ("em", "i"):
            _runs(par, child, bold, True)
        elif child.name == "br":
            par.add_run().add_break()
        else:
            _runs(par, child, bold, italic)


def html_to_docx(html: str) -> bytes:
    soup = BeautifulSoup(html or "", "html.parser")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    root = soup.body or soup
    for el in root.children:
        if isinstance(el, NavigableString):
            if str(el).strip():
                doc.add_paragraph(str(el).strip())
            continue
        name = el.name
        if name in ("h1", "h2", "h3"):
            doc.add_heading(el.get_text(" ", strip=True), level=int(name[1]))
        elif name in ("ol", "ul"):
            style = "List Number" if name == "ol" else "List Bullet"
            for li in el.find_all("li", recursive=False):
                p = doc.add_paragraph(style=style)
                _runs(p, li)
        elif name == "table":
            rows = el.find_all("tr")
            if rows:
                cols = max(len(r.find_all(["td", "th"])) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j, c in enumerate(r.find_all(["td", "th"])):
                        cell = table.cell(i, j)
                        cell.text = ""
                        _runs(cell.paragraphs[0], c, bold=(c.name == "th"))
        else:
            p = doc.add_paragraph()
            _runs(p, el)

    note = doc.add_paragraph()
    r = note.add_run(DISCLAIMER)
    r.italic = True
    r.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
